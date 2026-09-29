"""Generates realistic-looking (synthetic) CMPDI/CIL-style documents for the demo.
Creates: PDFs (narrative + table), an Excel reserves sheet, a DOCX parliamentary brief,
and ground_truth.csv (used to measure extraction accuracy)."""
import os, random
import numpy as np, pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
import docx

MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep"]
FULL = {"Apr": "April", "May": "May", "Jun": "June", "Jul": "July", "Aug": "August", "Sep": "September"}
SUBS = {"SECL": 15.0, "NCL": 12.0, "WCL": 5.5, "MCL": 20.0}
ISSUES = [
    "Progress of land acquisition for the expansion project was delayed by pending village rehabilitation.",
    "Railway rake availability constrained coal despatch during the month.",
    "Machinery breakdown of a dragline reduced overburden removal and output.",
    "Environmental clearance for capacity enhancement remains pending with the Ministry.",
    "A safety audit recommended fresh slope stability compliance under DGMS guidelines.",
    "Shortage of explosives supply affected blasting schedules at two opencast mines.",
    "Contractor manpower shortage slowed excavation and coal transportation.",
]
RAIN = "Heavy monsoon rainfall disrupted overburden removal, water pumping and internal haulage."


def generate(out_dir="data", seed=42):
    rng = random.Random(seed); np.random.seed(seed)
    os.makedirs(out_dir, exist_ok=True)
    styles = getSampleStyleSheet()
    truth = []
    for sub, base in SUBS.items():
        rows, narr = [], []
        for m in MONTHS:
            tgt = round(base * rng.uniform(0.98, 1.02), 2)
            act = round(tgt * rng.uniform(0.90, 1.05), 2)
            if sub == "WCL" and m == "Jul":
                act = round(tgt * 0.60, 2)                      # injected flood shortfall
            des = round(act * rng.uniform(0.97, 1.04), 2)
            ob = round(act * rng.uniform(2.2, 2.8), 2)
            ach = round(act / tgt * 100, 1)
            stated = ach
            if sub == "MCL" and m == "Aug":
                stated = round(ach - 10.0, 1)                   # injected typo in stated %
            narr_act = act
            if sub == "SECL" and m == "Sep":
                s = f"{act:.2f}"; narr_act = float(s[:-2] + s[-1] + s[-2])   # digit swap typo
            issue = RAIN if m in ("Jul", "Aug") else rng.choice(ISSUES)
            yr = 2025
            rows.append([f"{m} {yr}", f"{tgt:.2f}", f"{act:.2f}", f"{des:.2f}", f"{ob:.2f}", f"{stated:.1f}"])
            narr.append(f"In {FULL[m]} {yr}, {sub} produced {narr_act:.2f} MT against a target of {tgt:.2f} MT "
                        f"({ach:.1f}% achievement). {issue}")
            truth.append(dict(Subsidiary=sub, Month=f"{m} {yr}", Target_MT=tgt, Actual_MT=act,
                              Despatch_MT=des, OB_Removal_Mm3=ob))
        path = os.path.join(out_dir, f"Monthly_Production_Review_{sub}.pdf")
        doc = SimpleDocTemplate(path, pagesize=A4)
        el = [Paragraph(f"{sub} - Monthly Production Review, FY 2025-26", styles["Title"]),
              Paragraph("Prepared for the Ministry of Coal (synthetic demo data)", styles["Italic"]), Spacer(1, 12)]
        for n in narr:
            el += [Paragraph(n, styles["BodyText"]), Spacer(1, 6)]
        el.append(PageBreak())
        el.append(Paragraph(f"Annexure-I: Production statistics of {sub}", styles["Heading2"]))
        data = [["Month", "Target (MT)", "Actual (MT)", "Despatch (MT)", "OB Removal (Mm3)", "Achievement (%)"]] + rows
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
        el.append(t)
        doc.build(el)
    pd.DataFrame(truth).to_csv(os.path.join(out_dir, "ground_truth.csv"), index=False)

    # Geological reserves (Excel)
    blocks = []
    for sub in SUBS:
        for i in range(1, 4):
            blocks.append(dict(Subsidiary=sub, Block=f"{sub}-Block-{i}", Seam=rng.choice(["Seam-I", "Seam-IV", "Seam-VII"]),
                               Grade=rng.choice(["G7", "G9", "G11", "G13"]),
                               Proved_MT=round(rng.uniform(80, 900), 1), Indicated_MT=round(rng.uniform(20, 400), 1),
                               Avg_Depth_m=rng.randint(60, 420)))
    pd.DataFrame(blocks).to_excel(os.path.join(out_dir, "Geological_Reserves_Statement.xlsx"), index=False)

    # Parliamentary brief (Word)
    d = docx.Document()
    d.add_heading("Brief for Lok Sabha Unstarred Question - Coal Production and Land Acquisition", 1)
    for p in [
        "The Ministry has sought details on production shortfall and delays in land acquisition across coal companies.",
        "Land acquisition and rehabilitation of project affected families remain the principal cause of delay in expansion projects.",
        "Monsoon rainfall in July and August reduced overburden removal and coal despatch in western and central coalfields.",
        "Steps taken include daily monitoring of rake loading, faster forest clearance follow-up and deployment of additional dumpers.",
        "Safety compliance and DGMS audits are being tracked, with special attention to slope stability in opencast mines.",
    ]:
        d.add_paragraph(p)
    d.save(os.path.join(out_dir, "Parliament_Question_Brief.docx"))
    return out_dir


if __name__ == "__main__":
    print("Generated in", generate())
